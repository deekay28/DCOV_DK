#!/usr/bin/env python3
"""
build_seed_database.py
======================
ETL: converts the raw inspection worksheets (data/source/*.tsv, policy.json)
into the canonical component schema used by DCOV, and emits:

    data/components_seed.csv
    data/components_seed.json
    data/DCOV_Sample_Database.xlsx     <- import-wizard-ready workbook
    data/dcov_seed.sqlite              <- pre-built offline database
    web_demo/seed_data.js              <- embedded dataset for the offline web demo

Run:  python scripts/build_seed_database.py
"""
from __future__ import annotations

import csv
import hashlib
import json
import re
import sqlite3
import sys
from datetime import datetime, timezone
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "data" / "source"
OUT = ROOT / "data"

NOW = datetime.now(timezone.utc).strftime("%Y-%m-%d")

FIELDS = [
    "component_id", "component_name", "part_number", "chip_number", "manufacturer",
    "manufacturer_country", "country_of_origin", "is_chinese", "category",
    "drone_subsystem", "criticality", "criticality_policy", "alternative_manufacturer",
    "military_grade", "barcode", "qr_code", "function", "remarks", "image_path",
    "datasheet_url", "supplier", "date_added", "last_updated", "verified_by",
    "verification_source", "confidence_score",
]

POLICY = json.loads((SRC / "policy.json").read_text())
CHINESE_TOKENS = [t.upper() for t in POLICY["chinese_territories"]]
ALIASES = POLICY["manufacturer_aliases"]
ALTS = POLICY["approved_alternatives"]
CRIT = POLICY["criticality"]


# --------------------------------------------------------------------------- #
# Normalisation helpers (mirrored 1:1 in backend/app/services/normalize.py and
# in the Flutter client, so that offline and online matching agree.)
# --------------------------------------------------------------------------- #
def normalize_part(value: str) -> str:
    """Uppercase, strip everything that is not A-Z0-9. 'stm32 f302-c8t6' -> STM32F302C8T6"""
    return re.sub(r"[^A-Z0-9]", "", (value or "").upper())


def canonical_manufacturer(name: str) -> str:
    key = re.sub(r"[.,]", "", (name or "").strip()).upper()
    key = re.sub(r"\s+", " ", key)
    return ALIASES.get(key, key)


def is_chinese_origin(country: str, remarks: str = "") -> str:
    blob = f"{country} {remarks}".upper()
    if any(tok in blob for tok in ("MADE IN CHINA", "COO CHINA", "COO: CHINA")):
        return "YES"
    country_u = (country or "").upper().strip()
    if not country_u:
        return "UNKNOWN"
    if any(country_u == tok or tok in country_u.split("/") for tok in CHINESE_TOKENS):
        return "YES"
    if "CHINA" in country_u or "SHENZHEN" in country_u or "HONG KONG" in country_u:
        return "YES"
    return "NO"


def subsystem_of(item: str) -> str:
    """Map a free-text item name onto one of the policy subsystem buckets."""
    t = (item or "").upper()
    table = [
        ("FLIGHT CONTROL", "Flight Controller"), ("AUTO PILOT", "Flight Controller"),
        ("FCU", "Flight Controller"), ("FCS", "Flight Controller"),
        ("ON BOARD COMPUTER", "On Board Computer"), ("OBC", "On Board Computer"),
        ("JETSON", "On Board Computer"), ("COMPANION COMPUTER", "On Board Computer"),
        ("CRPA", "CRPA"), ("GNSS", "GNSS Receiver"), ("GPS", "GNSS Receiver"),
        ("RTK", "RTK GNSS Receiver"),
        ("RADIO", "Radio Modem"), ("MODEM", "Radio Modem"), ("GDT", "Radio Modem"),
        ("ADT", "Radio Modem"), ("TRANSCEIVER", "Radio Modem"),
        ("RECEIVER MODULE", "Radio Modem"), ("VIDEO RECEIVER", "Radio Modem"),
        ("GCS", "GCS"), ("GROUND CONTROL", "GCS"),
        ("ANTENNA", "Antenna"), ("JOYSTICK", "Joystick Controller"),
        ("ESC", "ESC"), ("SPEED CONTROLLER", "ESC"),
        ("MOTOR", "Motor"), ("PROPELLER", "Propeller"), ("ACTUATOR", "Propeller"),
        ("SERVO", "Propeller"),
        ("BATTERY", "Battery"), ("BTY", "Battery"), ("CHARGER", "Battery Charger"),
        ("POWER DISTR", "Power Distribution Board"), ("POWER MODULE", "Power Distribution Board"),
        ("POWER SUPPLY", "Power Distribution Board"), ("DC-DC", "Power Distribution Board"),
        ("VOLTAGE REG", "Power Distribution Board"),
        ("CAMERA", "Payload"), ("EO/IR", "Payload"), ("EO ", "Payload"),
        ("GIMBAL", "Payload"), ("STREAMING", "Payload"), ("ENCODER", "Payload"),
        ("HDMI", "Payload"), ("OFC", "Payload"),
        ("ETHERNET", "Ethernet Switch"), ("NETWORK SWITCH", "Ethernet Switch"),
        ("ROUTER", "Ethernet Switch"), ("SWITCH", "Ethernet Switch"),
        ("AIRFRAME", "Airframe"),
        ("SSD", "On Board Computer"), ("WIFI", "On Board Computer"),
        ("SENSOR", "Sensors"), ("IMU", "Sensors"), ("BARO", "Sensors"),
        ("AIR SPEED", "Sensors"), ("ENGINE CONTROL", "Propulsion"),
    ]
    for needle, bucket in table:
        if needle in t:
            return bucket
    return "Other"


def check_digit_ean13(body12: str) -> str:
    s = sum(int(d) * (3 if i % 2 else 1) for i, d in enumerate(body12))
    return str((10 - s % 10) % 10)


def synth_barcode(component_id: str) -> str:
    """Deterministic, checksum-valid EAN-13 for demo/label printing (890 = India GS1 prefix)."""
    digest = hashlib.sha1(component_id.encode()).hexdigest()
    body = "890" + "".join(c for c in digest if c.isdigit())[:9]
    body = (body + "000000000000")[:12]
    return body + check_digit_ean13(body)


def make_id(prefix: str, *parts: str) -> str:
    h = hashlib.sha1("|".join(parts).encode()).hexdigest()[:8].upper()
    return f"{prefix}-{h}"


# --------------------------------------------------------------------------- #
def read_tsv(name: str) -> list[dict]:
    with (SRC / name).open(newline="", encoding="utf-8") as fh:
        return list(csv.DictReader(fh, delimiter="\t"))


def blank_row() -> dict:
    r = {f: "" for f in FIELDS}
    r["military_grade"] = "NO"
    r["is_chinese"] = "UNKNOWN"
    r["date_added"] = NOW
    r["last_updated"] = NOW
    r["confidence_score"] = "100"
    return r


def enrich(row: dict) -> dict:
    sub = row["drone_subsystem"] or subsystem_of(row["component_name"])
    row["drone_subsystem"] = sub
    pol = CRIT.get(sub)
    if pol:
        row["criticality"] = pol["level"]
        row["criticality_policy"] = pol["policy"]
        if not row["remarks"]:
            row["remarks"] = pol["note"]
    else:
        row["criticality"] = row["criticality"] or "REVIEW"
        row["criticality_policy"] = row["criticality_policy"] or "Not classified in policy matrix - refer to inspecting authority"
    if not row["alternative_manufacturer"]:
        alt = ALTS.get(sub) or ALTS.get(row["category"])
        if alt:
            row["alternative_manufacturer"] = "; ".join(alt)
    row["manufacturer"] = canonical_manufacturer(row["manufacturer"])
    if not row["barcode"]:
        row["barcode"] = synth_barcode(row["component_id"])
    if not row["qr_code"]:
        row["qr_code"] = f"DCOV:{row['component_id']}:{normalize_part(row['chip_number'] or row['part_number'])}"
    if row["criticality"] == "CRITICAL":
        row["military_grade"] = "YES"
    return row


def build() -> list[dict]:
    rows: dict[str, dict] = {}

    def upsert(row: dict) -> None:
        cid = row["component_id"]
        if cid in rows:  # merge: keep the most specific origin verdict
            old = rows[cid]
            if old["is_chinese"] in ("UNKNOWN", "") and row["is_chinese"] != "UNKNOWN":
                rows[cid] = row
            return
        rows[cid] = row

    # ---- 1. KEY OBSNS: inspected units, chip-level findings -----------------
    for r in read_tsv("key_obsns.tsv"):
        chip = r["chip"].strip()
        if not chip:
            continue
        row = blank_row()
        row["chip_number"] = chip
        row["part_number"] = chip
        row["component_name"] = f"{r['item'].strip()} - {chip}"
        row["manufacturer"] = r["manufacturer"].strip() or "Unknown"
        verdict = is_chinese_origin("", r["remarks"])
        row["is_chinese"] = verdict
        row["country_of_origin"] = "China" if verdict == "YES" else "Unknown"
        row["manufacturer_country"] = ""
        row["category"] = "Chip / IC"
        row["drone_subsystem"] = subsystem_of(r["item"])
        row["function"] = r["function"].strip()
        row["supplier"] = r["model_make"].strip()
        row["remarks"] = r["remarks"].strip()
        row["verified_by"] = "Technical Inspection Team"
        row["verification_source"] = "Physical teardown - KEY OBSNS worksheet"
        row["confidence_score"] = "95" if verdict == "YES" else "70"
        row["component_id"] = make_id("CHIP", normalize_part(chip), row["manufacturer"].upper())
        upsert(enrich(row))

    # ---- 2. COMP IDEN AND ORIGIN: OEM + sub-component origins ---------------
    for r in read_tsv("comp_iden_origin.tsv"):
        sub = r["sub_component"].strip()
        oem = r["oem"].strip()
        country = r["country"].strip()
        main = r["main_component"].strip()
        row = blank_row()
        if sub:
            row["chip_number"] = sub
            row["part_number"] = sub
            row["component_name"] = f"{main} - {sub}"
            row["category"] = "Chip / IC"
            row["component_id"] = make_id("CHIP", normalize_part(sub), canonical_manufacturer(oem))
        else:
            row["part_number"] = main
            row["component_name"] = main
            row["category"] = "Assembly / Module"
            row["component_id"] = make_id("ASSY", main.upper(), canonical_manufacturer(oem))
        row["manufacturer"] = oem
        row["manufacturer_country"] = country
        row["country_of_origin"] = country
        row["is_chinese"] = is_chinese_origin(country)
        row["drone_subsystem"] = subsystem_of(main)
        row["verified_by"] = "Technical Inspection Team"
        row["verification_source"] = "Component identification worksheet (COMP IDEN AND ORIGIN)"
        row["confidence_score"] = "90"
        upsert(enrich(row))

    # ---- 3. COMMON CHINESE watchlist ---------------------------------------
    for r in read_tsv("common_chinese.tsv"):
        chip = r["chip_ref"].strip()
        row = blank_row()
        row["chip_number"] = chip
        row["part_number"] = chip
        row["component_name"] = f"Watchlist IC - {chip}"
        row["manufacturer"] = "Various (China-assembled)"
        row["manufacturer_country"] = "China"
        row["country_of_origin"] = "China"
        row["is_chinese"] = "YES"
        row["category"] = "Chip / IC"
        row["drone_subsystem"] = subsystem_of(r["subsystem"])
        row["remarks"] = f"On the standing Chinese-IC watchlist. Typically found in: {r['subsystem']}"
        row["verified_by"] = "Technical Inspection Team"
        row["verification_source"] = "COMMON CHINESE watchlist"
        row["confidence_score"] = "98"
        row["component_id"] = make_id("WATCH", normalize_part(chip))
        upsert(enrich(row))

    # ---- 4. OEM landscape ---------------------------------------------------
    for r in read_tsv("oem_landscape.tsv"):
        oem, country, cat = r["oem"].strip(), r["country"].strip(), r["category"].strip()
        row = blank_row()
        row["component_name"] = f"{cat} - {oem}"
        row["part_number"] = oem
        row["manufacturer"] = oem
        row["manufacturer_country"] = country
        row["country_of_origin"] = country
        row["is_chinese"] = is_chinese_origin(country)
        row["category"] = "OEM / LRU"
        row["drone_subsystem"] = subsystem_of(cat) if subsystem_of(cat) != "Other" else cat
        row["verified_by"] = "Procurement Cell"
        row["verification_source"] = "OEM landscape (DRONE AND LM PARTS)"
        row["confidence_score"] = "85"
        row["component_id"] = make_id("OEM", cat.upper(), canonical_manufacturer(oem))
        upsert(enrich(row))

    return sorted(rows.values(), key=lambda r: (r["category"], r["component_name"]))


# --------------------------------------------------------------------------- #
def write_csv(rows: list[dict]) -> None:
    p = OUT / "components_seed.csv"
    with p.open("w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=FIELDS)
        w.writeheader()
        w.writerows(rows)
    print(f"  csv    -> {p.relative_to(ROOT)}  ({len(rows)} rows)")


def write_json(rows: list[dict]) -> None:
    p = OUT / "components_seed.json"
    p.write_text(json.dumps(rows, indent=1))
    print(f"  json   -> {p.relative_to(ROOT)}")


def write_sqlite(rows: list[dict]) -> None:
    p = OUT / "dcov_seed.sqlite"
    if p.exists():
        p.unlink()
    con = sqlite3.connect(p)
    cols = ", ".join(f'"{f}" TEXT' for f in FIELDS)
    con.execute(f"CREATE TABLE components (id INTEGER PRIMARY KEY AUTOINCREMENT, {cols})")
    con.execute("CREATE UNIQUE INDEX ix_components_cid ON components(component_id)")
    con.execute("CREATE INDEX ix_components_chip ON components(chip_number)")
    con.execute("CREATE INDEX ix_components_barcode ON components(barcode)")
    con.execute("CREATE INDEX ix_components_mfr ON components(manufacturer)")
    # normalised search key column, used by the exact-match fast path
    con.execute("ALTER TABLE components ADD COLUMN search_key TEXT")
    con.execute("CREATE INDEX ix_components_key ON components(search_key)")
    ph = ", ".join("?" for _ in FIELDS) + ", ?"
    con.executemany(
        f'INSERT INTO components ({", ".join(chr(34)+f+chr(34) for f in FIELDS)}, search_key) VALUES ({ph})',
        [tuple(r[f] for f in FIELDS) + (normalize_part(r["chip_number"] or r["part_number"]),) for r in rows],
    )
    # FTS5 index for keyword search
    con.execute("CREATE VIRTUAL TABLE components_fts USING fts5(component_name, part_number, chip_number, manufacturer, remarks, content='')")
    con.executemany(
        "INSERT INTO components_fts (rowid, component_name, part_number, chip_number, manufacturer, remarks) VALUES (?,?,?,?,?,?)",
        [(i + 1, r["component_name"], r["part_number"], r["chip_number"], r["manufacturer"], r["remarks"]) for i, r in enumerate(rows)],
    )
    con.commit()
    con.close()
    print(f"  sqlite -> {p.relative_to(ROOT)}")


def write_xlsx(rows: list[dict]) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
        from openpyxl.worksheet.table import Table, TableStyleInfo
    except ImportError:
        print("  xlsx   -> skipped (openpyxl not installed)")
        return

    wb = Workbook()
    ws = wb.active
    ws.title = "COMPONENTS"
    header_fill = PatternFill("solid", fgColor="1F3864")
    ws.append([f.replace("_", " ").upper() for f in FIELDS])
    for c in ws[1]:
        c.font = Font(name="Arial", bold=True, color="FFFFFF", size=10)
        c.fill = header_fill
        c.alignment = Alignment(vertical="center", wrap_text=True)
    red = PatternFill("solid", fgColor="F8CBAD")
    yellow = PatternFill("solid", fgColor="FFE699")
    for r in rows:
        ws.append([r[f] for f in FIELDS])
        row_cells = ws[ws.max_row]
        for c in row_cells:
            c.font = Font(name="Arial", size=10)
        if r["is_chinese"] == "YES":
            for c in row_cells:
                c.fill = red
        elif r["is_chinese"] == "UNKNOWN":
            for c in row_cells:
                c.fill = yellow
    widths = {"component_name": 42, "remarks": 60, "alternative_manufacturer": 55,
              "criticality_policy": 45, "verification_source": 40, "component_id": 16,
              "chip_number": 20, "part_number": 22, "manufacturer": 26, "qr_code": 28}
    for i, f in enumerate(FIELDS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(f, 16)
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(FIELDS))}{ws.max_row}"

    # ---- LEGEND sheet -------------------------------------------------------
    lg = wb.create_sheet("LEGEND")
    lg.append(["DCOV IMPORT TEMPLATE - FIELD LEGEND"])
    lg["A1"].font = Font(name="Arial", bold=True, size=14)
    lg.append([])
    lg.append(["FIELD", "REQUIRED", "MEANING / ACCEPTED VALUES"])
    for c in lg[3]:
        c.font = Font(name="Arial", bold=True, color="FFFFFF")
        c.fill = header_fill
    legend = [
        ("component_id", "Auto", "Left blank on import, DCOV generates a stable ID from part number + manufacturer."),
        ("component_name", "YES", "Human-readable description shown on the result banner."),
        ("part_number", "YES*", "Manufacturer part number. *Either part_number or chip_number must be present."),
        ("chip_number", "YES*", "Marking as laser-etched/printed on the die package - what OCR reads."),
        ("manufacturer", "YES", "Silicon vendor or OEM. Aliases are resolved automatically (e.g. 'ST Micro' -> STMICROELECTRONICS)."),
        ("manufacturer_country", "No", "Country of the vendor's corporate HQ."),
        ("country_of_origin", "YES", "Assembly/test country - this drives the verdict, not the HQ country."),
        ("is_chinese", "Auto", "YES / NO / UNKNOWN. Recomputed on import from country_of_origin and remarks."),
        ("category", "No", "Chip / IC, Assembly / Module, OEM / LRU."),
        ("drone_subsystem", "No", "Mapped to the policy matrix: Flight Controller, GNSS Receiver, CRPA, Radio Modem, GCS, ESC, ..."),
        ("criticality", "Auto", "CRITICAL / NON-CRITICAL, derived from drone_subsystem via the policy matrix."),
        ("alternative_manufacturer", "No", "Approved substitutes (seeded from the NDDA / Blue UAS list)."),
        ("military_grade", "No", "YES / NO."),
        ("barcode", "No", "If blank, a checksum-valid EAN-13 is generated for label printing."),
        ("qr_code", "No", "If blank, generated as DCOV:<component_id>:<normalised part>."),
        ("remarks", "No", "Free text. Phrases 'COO China' / 'Made in China' force is_chinese = YES."),
        ("supplier", "No", "Where the unit was procured from / the parent assembly it was pulled out of."),
        ("verified_by", "No", "Name or cell that confirmed the origin."),
        ("verification_source", "No", "Evidence trail: teardown report, datasheet, vendor letter, customs declaration."),
        ("confidence_score", "No", "0-100. Below 60 the result screen is downgraded to NEEDS VERIFICATION."),
    ]
    for f, req, meaning in legend:
        lg.append([f, req, meaning])
    for row in lg.iter_rows(min_row=4):
        for c in row:
            c.font = Font(name="Arial", size=10)
            c.alignment = Alignment(vertical="top", wrap_text=True)
    lg.column_dimensions["A"].width = 26
    lg.column_dimensions["B"].width = 12
    lg.column_dimensions["C"].width = 95
    lg.append([])
    lg.append(["EXAMPLE ROW (delete before import):"])
    lg.append(["", "component_name", "GNSS Receiver - STM32G4A1KCU6"])
    lg.append(["", "chip_number", "STM32G4A1KCU6"])
    lg.append(["", "manufacturer", "STMicroelectronics"])
    lg.append(["", "country_of_origin", "China"])
    lg.append(["", "drone_subsystem", "GNSS Receiver"])
    lg.append(["", "remarks", "COO China - marking CHN on package underside"])
    lg.append([])
    lg.append(["Source: Chinese_Components_FINAL.xlsx (sheets DRONE AND LM PARTS, COMP IDEN AND ORIGIN,"])
    lg.append(["KEY OBSNS, COMMON CHINESE, CRITICAL NON CRITICAL COMPONENT, NDDA & BLUE UAS LIST),"])
    lg.append([f"normalised by scripts/build_seed_database.py on {NOW}."])

    # ---- SUMMARY sheet (formulas, not hardcoded) ----------------------------
    sm = wb.create_sheet("SUMMARY", 0)
    sm.append(["DCOV SEED DATABASE - SUMMARY"])
    sm["A1"].font = Font(name="Arial", bold=True, size=14)
    sm.append([])
    n = len(rows)
    ic_col = get_column_letter(FIELDS.index("is_chinese") + 1)
    cr_col = get_column_letter(FIELDS.index("criticality") + 1)
    rng_ic = f"COMPONENTS!${ic_col}$2:${ic_col}${n + 1}"
    rng_cr = f"COMPONENTS!${cr_col}$2:${cr_col}${n + 1}"
    sm.append(["Metric", "Count"])
    for c in sm[3]:
        c.font = Font(name="Arial", bold=True, color="FFFFFF")
        c.fill = header_fill
    sm.append(["Total components", f"=COUNTA({rng_ic})"])
    sm.append(["Chinese origin", f'=COUNTIF({rng_ic},"YES")'])
    sm.append(["Non-Chinese origin", f'=COUNTIF({rng_ic},"NO")'])
    sm.append(["Unknown origin", f'=COUNTIF({rng_ic},"UNKNOWN")'])
    sm.append(["Critical subsystems", f'=COUNTIF({rng_cr},"CRITICAL")'])
    sm.append(["Chinese AND critical (escalate)", f'=COUNTIFS({rng_ic},"YES",{rng_cr},"CRITICAL")'])
    for row in sm.iter_rows(min_row=4):
        for c in row:
            c.font = Font(name="Arial", size=10)
    sm.column_dimensions["A"].width = 34
    sm.column_dimensions["B"].width = 12
    sm.append([])
    sm.append(["All counts are live formulas over the COMPONENTS sheet - they update if you edit rows."])

    p = OUT / "DCOV_Sample_Database.xlsx"
    wb.save(p)
    print(f"  xlsx   -> {p.relative_to(ROOT)}")


def write_web_seed(rows: list[dict]) -> None:
    p = ROOT / "web_demo" / "seed_data.js"
    payload = {
        "generated": NOW,
        "source": "Chinese_Components_FINAL.xlsx",
        "fields": FIELDS,
        "policy": {"criticality": CRIT, "coo_marking_codes": POLICY["coo_marking_codes"],
                   "aliases": ALIASES},
        "components": rows,
    }
    p.write_text("// AUTO-GENERATED by scripts/build_seed_database.py - do not edit.\n"
                 "window.DCOV_SEED = " + json.dumps(payload, separators=(",", ":")) + ";\n")
    print(f"  webjs  -> {p.relative_to(ROOT)}")


def main() -> int:
    print("Building DCOV seed database...")
    rows = build()
    write_csv(rows)
    write_json(rows)
    write_sqlite(rows)
    write_xlsx(rows)
    write_web_seed(rows)
    chinese = sum(1 for r in rows if r["is_chinese"] == "YES")
    unknown = sum(1 for r in rows if r["is_chinese"] == "UNKNOWN")
    crit_cn = sum(1 for r in rows if r["is_chinese"] == "YES" and r["criticality"] == "CRITICAL")
    print(f"\n  total={len(rows)}  chinese={chinese}  non-chinese={len(rows) - chinese - unknown}  "
          f"unknown={unknown}  chinese+critical={crit_cn}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
